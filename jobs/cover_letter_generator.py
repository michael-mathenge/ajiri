from io import BytesIO
from datetime import date
from docx import Document
from .matching import calculate_match_score


def generate_cover_letter_docx(profile, job):
    """
    Builds a tailored cover letter as a .docx, referencing both the
    candidate's profile AND the specific job being applied to. Template-based
    (not AI-generated) for v1 — predictable, free, no external API dependency.
    """
    doc = Document()

    name = profile.user.full_name or profile.user.email.split('@')[0].title()
    match_score = calculate_match_score(profile, job)

    header = doc.add_paragraph()
    header.add_run(name).bold = True
    if profile.location:
        doc.add_paragraph(profile.location)
    doc.add_paragraph(profile.user.email)
    doc.add_paragraph(date.today().strftime('%B %d, %Y'))

    doc.add_paragraph()

    doc.add_paragraph(f"Dear Hiring Manager at {job.company_name},")
    doc.add_paragraph()

    # Opening line tone shifts based on how well the profile's skills match
    # this specific job's requirements, using the same calculate_match_score()
    # logic that powers the Discovery browse/search results.
    if match_score >= 75:
        opening = (
            f"I am excited to apply for the {job.title} position at {job.company_name} — "
            f"my background is a strong match for what you're looking for. "
        )
    elif match_score >= 40:
        opening = (
            f"I am writing to express my interest in the {job.title} position at "
            f"{job.company_name}, where I believe my skills would be a solid contribution. "
        )
    else:
        opening = (
            f"I am writing to express my interest in the {job.title} position at "
            f"{job.company_name}. "
        )

    if profile.headline:
        opening += f"As a {profile.headline}, I believe my background aligns well with this role."
    doc.add_paragraph(opening)

    shared_skills = set(profile.skills.values_list('name', flat=True)) & set(
        job.skills_required.values_list('name', flat=True)
    )
    body_parts = []
    if profile.years_of_experience:
        body_parts.append(
            f"I bring {profile.years_of_experience} year(s) of relevant professional experience."
        )
    if shared_skills:
        skills_list = ', '.join(sorted(shared_skills))
        body_parts.append(
            f"My experience with {skills_list} directly matches the requirements for this role."
        )
    if profile.bio:
        body_parts.append(profile.bio)

    if body_parts:
        doc.add_paragraph(' '.join(body_parts))

    doc.add_paragraph()
    closing = (
        f"I would welcome the opportunity to discuss how my skills can contribute "
        f"to {job.company_name}. Thank you for considering my application."
    )
    doc.add_paragraph(closing)

    doc.add_paragraph()
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph(name)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer