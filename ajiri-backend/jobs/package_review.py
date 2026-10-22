"""Verify saved bytes before preview, download or explicit document approval.

see docs/CONCEPTS.md#application-package-approval
"""
import hashlib
from io import BytesIO

from docx import Document
from django.db import OperationalError, transaction
from django.db.models import F
from django.shortcuts import get_object_or_404

from profiles.cv_builder import content_hash, CVUnavailable
from profiles.cv_validation import MAX_DOCUMENT_BYTES
from profiles.upload_validation import _docx
from .application_packages import package_data, package_fingerprint, require_editable, is_editable, PackageConflict
from .models import Application, ApplicationPackage, ApplicationPackageApproval

POLICY_VERSION = 'application-documents-v1'
APPROVAL_STATEMENT = ('I reviewed the saved CV and cover letter for this revision and approve these '
                      'documents only. This does not authorize sending or submission.')


class PackageUnavailable(CVUnavailable):
    default_detail = 'The saved document package is unavailable. Reload its revisions before continuing.'


def owner_package(user, application_id, package_id):
    return get_object_or_404(ApplicationPackage.objects.select_related(
        'application__job', 'cv_approval__profile'), pk=package_id,
        application_id=application_id, application__user=user)


def verified_documents(package):
    """Return the actual saved bytes, never regenerated files or storage URLs."""
    source = package.cv_approval
    if (source.profile.user_id != package.application.user_id
            or content_hash(source.content) != source.content_sha256
            or package.cv_sha256 != source.document_sha256
            or package_fingerprint(package) != package.content_sha256):
        raise PackageUnavailable()
    result = {}
    for kind, field, expected_size, digest in [
        ('cv', package.cv_document, package.cv_size, package.cv_sha256),
        ('cover-letter', package.letter_document, package.letter_size, package.letter_sha256),
    ]:
        try:
            if not 0 < expected_size <= MAX_DOCUMENT_BYTES:
                raise ValueError('Invalid size')
            with field.open('rb') as stream:
                data = stream.read(MAX_DOCUMENT_BYTES + 1)
            if len(data) != expected_size or hashlib.sha256(data).hexdigest() != digest:
                raise ValueError('Invalid bytes')
            # Reuse existing bounded package checks before any DOCX parser expansion.
            _docx(data)
        except Exception:
            raise PackageUnavailable() from None
        result[kind] = data
    return result


def document_preview(data):
    """Paragraph text/order from stored DOCX bytes; not a pixel/page-layout renderer."""
    try:
        document = Document(BytesIO(data))
        # Current Ajiri generators emit paragraphs only. Never silently omit an unexpected table.
        if document.tables or len(document.paragraphs) > 2000:
            raise ValueError('Unsupported layout')
        paragraphs = [{'text': row.text, 'style': row.style.name if row.style else 'Normal'}
                      for row in document.paragraphs]
        if sum(len(row['text']) for row in paragraphs) > 120000:
            raise ValueError('Oversized preview')
        return paragraphs
    except Exception:
        raise PackageUnavailable() from None


def approval_data(package):
    proof = ApplicationPackageApproval.objects.filter(package=package).first()
    if not proof:
        return None
    if (proof.approved_by_id != package.application.user_id
            or proof.content_sha256 != package.content_sha256
            or proof.policy_version != POLICY_VERSION):
        raise PackageUnavailable()
    return {'approved_at': proof.approved_at, 'policy_version': proof.policy_version,
            'content_sha256': proof.content_sha256}


def review_data(package):
    documents = verified_documents(package)
    previews = {kind: document_preview(data) for kind, data in documents.items()}
    if '\n'.join(row['text'] for row in previews['cover-letter']) != package.letter_text.replace('\r', '\n'):
        raise PackageUnavailable()
    app = package.application
    latest = app.packages.order_by('-revision').values_list('revision', flat=True).first()
    return {**package_data(package), 'previews': previews,
        'approval': approval_data(package), 'is_latest': package.revision == latest,
        'approvable': package.revision == latest and is_editable(app),
        'approval_statement': APPROVAL_STATEMENT, 'policy_version': POLICY_VERSION}


def approve_package(user, application_id, package_id, expected_revision, expected_hash):
    # Scope the lookup before attempting a lock, preserving 404 for cross-owner requests.
    owner_package(user, application_id, package_id)
    try:
        with transaction.atomic():
            if not Application.objects.filter(pk=application_id, user=user,
                    status__in=['draft', 'ready_for_review']).update(status=F('status')):
                raise PackageConflict()
            package = owner_package(user, application_id, package_id)
            require_editable(package.application)
            latest = package.application.packages.order_by('-revision').values_list('revision', flat=True).first()
            if (expected_revision != package.revision or latest != expected_revision
                    or package.content_sha256 != expected_hash):
                raise PackageConflict()
            # All checks run under the application writer claim; newer saves cannot interleave.
            # A filesystem can change later: every subsequent download/approval rechecks hashes.
            review_data(package)
            proof, created = ApplicationPackageApproval.objects.get_or_create(package=package,
                defaults={'approved_by': user, 'content_sha256': expected_hash, 'policy_version': POLICY_VERSION})
            return approval_data(package), created
    except OperationalError as error:
        if 'locked' in str(error).lower():
            raise PackageConflict('The application is busy. Reload its approval before continuing.') from None
        raise
