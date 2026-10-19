"""Synthetic structure checks; visual pagination is verified separately in Word."""
from io import BytesIO
from django.test import SimpleTestCase
from docx import Document
from docx.oxml.ns import qn
from .cv_builder import render_docx, TEMPLATE_VERSION
from .cv_validation import validate_content


def layout_fixture(long=False):
    return validate_content({'schema_version': 1,
        'contact': {'name': 'Synthetic Applicant With Several Name Parts', 'email': 'synthetic@example.invalid'},
        'summary': 'Synthetic layout review. ' * (30 if long else 1),
        'employment': [{'title': 'Senior engineering and programme delivery lead ' * (4 if long else 1),
            'organisation': 'Synthetic organisation with a long descriptive company name ' * (3 if long else 1),
            'start': '2000-01', 'end': '2026-09',
            'description': 'Synthetic paragraph for layout verification. ' * (16 if long else 1),
            'achievements': ['Synthetic achievement with readable wrapping. ' * (6 if long else 1)] * 3}
            for _ in range(7 if long else 1)],
        'education': [{'qualification': 'Synthetic computing qualification', 'institution': 'Synthetic college',
                       'start': '1995-01', 'end': '1999-12'}]}, for_approval=True)


class CVLayoutTests(SimpleTestCase):
    def test_resolved_styles_and_no_theme_border_or_implicit_tabs(self):
        document = Document(BytesIO(render_docx(layout_fixture(long=True))))
        self.assertEqual(TEMPLATE_VERSION, 'cv-v2')
        for name, size in [('Normal', 11), ('Title', 22), ('Heading 1', 14), ('Heading 2', 12), ('List Bullet', 11)]:
            style = document.styles[name]
            self.assertEqual(style.font.name, 'Calibri')
            self.assertEqual(style.font.size.pt, size)
            self.assertFalse(any('theme' in key.lower() for key in style.element.rPr.rFonts.attrib))
        self.assertFalse(document.styles['Title'].element.xpath('./w:pPr/w:pBdr'))
        self.assertFalse(document.tables)
        self.assertFalse(document.element.xpath('.//w:tab'))
        for paragraph in document.paragraphs:
            self.assertNotIn('\t', paragraph.text)
            if paragraph.style.name in ('Heading 1', 'Heading 2'):
                self.assertTrue(paragraph.style.paragraph_format.keep_with_next)
        for section in document.sections:
            self.assertGreater(section.page_width - section.left_margin - section.right_margin, 0)
            self.assertGreater(section.page_height - section.top_margin - section.bottom_margin, 0)

    def test_dates_are_separate_from_long_titles_and_lists_have_defined_numbering(self):
        document = Document(BytesIO(render_docx(layout_fixture(long=True))))
        paragraphs = document.paragraphs
        dates = [i for i, p in enumerate(paragraphs) if p.text == '2000-01 – 2026-09']
        self.assertEqual(len(dates), 7)
        for index in dates:
            self.assertEqual(paragraphs[index].style.name, 'Normal')
            self.assertEqual(paragraphs[index - 1].style.name, 'Heading 2')
            self.assertNotIn('2000-01', paragraphs[index - 1].text)
        numbering = document.part.numbering_part.element
        num_id = document.styles['List Bullet'].element.xpath('./w:pPr/w:numPr/w:numId')[0].get(qn('w:val'))
        self.assertTrue(numbering.xpath(f'./w:num[@w:numId="{num_id}"]'))
        self.assertTrue(any(p.style.name == 'List Bullet' for p in paragraphs))

    def test_optional_sections_omitted_and_plain_text_preserved(self):
        content = validate_content({'schema_version': 1, 'contact': {'name': 'Synthetic Applicant'},
            'summary': 'Reviewed literal <text> & details.'}, for_approval=True)
        document = Document(BytesIO(render_docx(content)))
        self.assertEqual([p.text for p in document.paragraphs],
            ['Synthetic Applicant', 'Summary', 'Reviewed literal <text> & details.'])
