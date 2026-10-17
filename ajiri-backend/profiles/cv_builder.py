"""Versioned personal CV workflow, independent of application snapshots.

see docs/CONCEPTS.md#cv-builder-flow
"""
from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import json
from docx import Document
from docx.shared import Pt
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.db.models import F
from django.utils import timezone
from rest_framework.exceptions import APIException
from .cv_validation import validate_content, MAX_DOCUMENT_BYTES
from .models import CVWorkspace, CVApproval, approved_cv_path

TEMPLATE_VERSION = 'cv-v1'


class CVConflict(APIException):
    status_code = 409
    default_detail = 'The CV changed. Reload it and review the current revision before continuing.'


class CVUnavailable(APIException):
    status_code = 503
    default_detail = 'The CV is unavailable. Please try again later.'


def content_hash(content):
    return hashlib.sha256(json.dumps(content, ensure_ascii=False, sort_keys=True,
                                     separators=(',', ':'), allow_nan=False).encode('utf-8')).hexdigest()


def layout(content):
    """The HTML preview and DOCX consume exactly this ordered plain-text layout."""
    contact = content['contact']
    result = {'name': contact['name'], 'headline': content['headline'],
              'contact': ' | '.join(contact[key] for key in ('location', 'email', 'phone') if contact[key]),
              'sections': []}
    sections = result['sections']
    if content['summary']:
        sections.append({'title': 'Summary', 'entries': [{'heading': '', 'paragraphs': [content['summary']], 'bullets': []}]})
    if content['skills']:
        sections.append({'title': 'Skills', 'entries': [{'heading': '', 'paragraphs': [', '.join(content['skills'])], 'bullets': []}]})
    for key, title in [('employment', 'Experience'), ('education', 'Education'),
                       ('certifications', 'Certifications'), ('projects', 'Projects')]:
        entries = []
        for row in content[key]:
            paragraphs = []
            if key in ('employment', 'education'):
                heading = ' — '.join(row[name] for name in (
                    ('title', 'organisation') if key == 'employment' else ('qualification', 'institution')))
                dates = [row['start'], 'Present' if row['current'] else row['end']]
                if any(dates):
                    paragraphs.append(' – '.join(date for date in dates if date))
            elif key == 'certifications':
                heading = f"{row['name']} — {row['issuer']}"
                paragraphs.extend(f'{label}: {row[name]}' for name, label in [('issued', 'Issued'), ('expiry', 'Expires')] if row[name])
            else:
                heading = ' — '.join(value for value in (row['name'], row['organisation']) if value)
                if row['url']:
                    paragraphs.append(row['url'])
            if row['description']:
                paragraphs.append(row['description'])
            entries.append({'heading': heading, 'paragraphs': paragraphs, 'bullets': row.get('achievements', [])})
        if entries:
            sections.append({'title': title, 'entries': entries})
    return result


def render_docx(content):
    preview = layout(content)
    document = Document()
    document.styles['Normal'].font.name = 'Calibri'
    document.styles['Normal'].font.size = Pt(11)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.created = datetime(2000, 1, 1)
    document.core_properties.modified = datetime(2000, 1, 1)
    document.add_heading(preview['name'], 0)
    for value in (preview['headline'], preview['contact']):
        if value:
            document.add_paragraph(value)
    for section in preview['sections']:
        document.add_heading(section['title'], 1)
        for entry in section['entries']:
            if entry['heading']:
                document.add_heading(entry['heading'], 2)
            for paragraph in entry['paragraphs']:
                document.add_paragraph(paragraph)
            for bullet in entry['bullets']:
                document.add_paragraph(bullet, style='List Bullet')
    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    if len(data) > MAX_DOCUMENT_BYTES:
        raise CVUnavailable()
    return data


def start_workspace(profile):
    existing = CVWorkspace.objects.filter(profile=profile).first()
    if existing:
        return existing
    seed = validate_content({'schema_version': 1,
        'contact': {'name': profile.user.full_name or '', 'email': profile.user.email,
                    'phone': profile.user.phone_number or '', 'location': profile.location},
        'headline': profile.headline, 'summary': profile.bio,
        'skills': list(profile.skills.values_list('name', flat=True))})
    return CVWorkspace.objects.get_or_create(profile=profile, defaults={'content': seed})[0]


def save_workspace(workspace, expected_revision, patch):
    if workspace.revision != expected_revision:
        raise CVConflict()
    merged = deepcopy(workspace.content)
    merged.update(patch)
    if isinstance(patch.get('contact'), dict):
        merged['contact'] = {**workspace.content.get('contact', {}), **patch['contact']}
    content = validate_content(merged)
    updated = timezone.now()
    changed = CVWorkspace.objects.filter(pk=workspace.pk, revision=expected_revision).update(
        content=content, revision=F('revision') + 1, updated_at=updated)
    if changed != 1:
        raise CVConflict()
    workspace.content, workspace.revision, workspace.updated_at = content, expected_revision + 1, updated
    return workspace


def approve_workspace(workspace, expected_revision, expected_hash):
    if workspace.revision != expected_revision or content_hash(workspace.content) != expected_hash:
        raise CVConflict()
    content = validate_content(workspace.content, for_approval=True)
    existing = CVApproval.objects.filter(profile=workspace.profile, revision=expected_revision).first()
    if existing:
        return existing
    data = render_docx(content)
    path = None
    try:
        path = default_storage.save(approved_cv_path(None, ''), ContentFile(data))
        with transaction.atomic():
            # Conditional write acquires SQLite's writer lock. A prior SELECT /
            # select_for_update is not enough. Concurrent edits lose this CAS.
            matched = CVWorkspace.objects.filter(pk=workspace.pk, revision=expected_revision).update(revision=F('revision'))
            if matched != 1:
                raise CVConflict()
            current = CVWorkspace.objects.get(pk=workspace.pk)
            if content_hash(current.content) != expected_hash:
                raise CVConflict()
            existing = CVApproval.objects.filter(profile=workspace.profile, revision=expected_revision).first()
            if existing:
                approval = existing
            else:
                approval = CVApproval.objects.create(profile=workspace.profile, revision=expected_revision,
                    content=content, content_sha256=expected_hash, template_version=TEMPLATE_VERSION,
                    document=path, document_sha256=hashlib.sha256(data).hexdigest())
        if existing:
            default_storage.delete(path)
        return approval
    except Exception:
        if path:
            # Only this request's uncommitted, uniquely generated file is removed.
            # A hard process crash can leave an orphan for report-only review.
            default_storage.delete(path)
        raise


def document_bytes(approval):
    try:
        with approval.document.open('rb') as stream:
            data = stream.read(MAX_DOCUMENT_BYTES + 1)
        if len(data) > MAX_DOCUMENT_BYTES or hashlib.sha256(data).hexdigest() != approval.document_sha256:
            raise CVUnavailable()
        return data
    except (OSError, ValueError):
        raise CVUnavailable() from None
