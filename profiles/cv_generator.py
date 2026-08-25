from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_cv_docx(profile):
    """
    Builds a clean, standardized CV as a .docx file from structured
    Profile data, rather than trying to parse/reformat an uploaded file
    (PDF/DOCX text extraction is unreliable across different layouts).
    Returns an in-memory file (BytesIO), ready to save or serve directly.
    """
    doc = Document()

    name = profile.user.email.split('@')[0].replace('.', ' ').title()
    title = doc.add_heading(name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if profile.headline:
        subtitle = doc.add_paragraph(profile.headline)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].italic = True

    contact_line = f"{profile.location or ''}  |  {profile.user.email}"
    contact = doc.add_paragraph(contact_line)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    if profile.bio:
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(profile.bio)

    doc.add_heading('Experience', level=1)
    doc.add_paragraph(f"{profile.years_of_experience} year(s) of professional experience")

    skills = profile.skills.all()
    if skills:
        doc.add_heading('Skills', level=1)
        skills_paragraph = doc.add_paragraph()
        skills_paragraph.add_run(', '.join(skill.name for skill in skills))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer